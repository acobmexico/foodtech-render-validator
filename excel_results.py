import csv
import os
import re
import threading
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CSV_PATH = Path(os.getenv("RESULTS_CSV_PATH", "/tmp/FOODTECH_VALIDACION_RESULTADOS.csv"))
XLSX_PATH = Path(os.getenv("RESULTS_XLSX_PATH", "/tmp/FOODTECH_VALIDACION_RESULTADOS.xlsx"))
_lock = threading.Lock()
_illegal = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")

HEADERS = [
    "FECHA_UTC", "ID", "ENROLLMENT_CODE", "ID_VISITANTE", "EMPRESA",
    "PAGINA_WEB", "ESTADO_PROCESO", "PUNTAJE", "TIPO", "DECISION",
    "GIRO_DETECTADO", "EVIDENCIA", "RAZONAMIENTO", "URL_PRINCIPAL",
    "URL_CORPORATIVA", "ERROR_WEB", "ERROR_PROCESO", "TIEMPO_SEGUNDOS",
    "MODELO", "METODO_OBTENCION", "CACHE_UTILIZADO", "INTENTOS_RECUPERACION",
    "FUENTES_BUSQUEDA",
]


def clean_cell(value):
    if value is None:
        return ""
    text = _illegal.sub("", str(value))[:32767]
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def result_row(result: dict, source_record: dict) -> list[str]:
    evaluation = result.get("evaluation") or {}
    return [clean_cell(value) for value in [
        evaluation.get("evaluated_at_utc") or datetime.now(timezone.utc).isoformat(),
        result.get("id", ""),
        result.get("enrollmentCode", ""), result.get("IDVisitante", ""),
        result.get("RazonSocial", ""), source_record.get("PaginaWeb", ""),
        result.get("status", ""), evaluation.get("score", ""),
        evaluation.get("type", ""), evaluation.get("decision", ""),
        evaluation.get("detected_business", ""), evaluation.get("evidence", ""),
        evaluation.get("reason", ""), evaluation.get("main_url", ""),
        evaluation.get("corporate_url", ""), evaluation.get("web_error", ""),
        result.get("error", ""), result.get("elapsed_seconds", ""),
        evaluation.get("model", ""), evaluation.get("retrieval_method", ""),
        "SI" if evaluation.get("cache_hit") else "NO",
        evaluation.get("retrieval_attempts", ""),
        " | ".join(evaluation.get("search_sources") or []),
    ]]


def _ensure_current_schema() -> None:
    if not CSV_PATH.exists() or CSV_PATH.stat().st_size == 0:
        return
    with CSV_PATH.open("r", newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    if not rows or rows[0] == HEADERS:
        return
    old_headers = rows[0]
    positions = {name: index for index, name in enumerate(old_headers)}
    temporary = CSV_PATH.with_suffix(".migration.tmp")
    with temporary.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        for old_row in rows[1:]:
            writer.writerow([
                old_row[positions[name]] if name in positions and positions[name] < len(old_row) else ""
                for name in HEADERS
            ])
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary, CSV_PATH)


def append_analysis_result(result: dict, source_record: dict) -> None:
    row = result_row(result, source_record)
    with _lock:
        CSV_PATH.parent.mkdir(parents=True, exist_ok=True)
        _ensure_current_schema()
        new_file = not CSV_PATH.exists() or CSV_PATH.stat().st_size == 0
        with CSV_PATH.open("a", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            if new_file:
                writer.writerow(HEADERS)
            writer.writerow(row)
            handle.flush()
            os.fsync(handle.fileno())


def build_excel() -> Path:
    with _lock:
        if not CSV_PATH.exists() or CSV_PATH.stat().st_size == 0:
            raise FileNotFoundError("Todavía no existen resultados para descargar.")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "RESULTADOS"
        with CSV_PATH.open("r", newline="", encoding="utf-8-sig") as handle:
            for row_number, row in enumerate(csv.reader(handle), start=1):
                if row_number > 1:
                    if len(row) >= 8 and row[7].isdigit():
                        row[7] = int(row[7])
                    if len(row) >= 18:
                        try:
                            row[17] = float(row[17])
                        except (TypeError, ValueError):
                            pass
                sheet.append(row)

        fill = PatternFill("solid", fgColor="174B32")
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = {
            1: 25, 2: 12, 3: 20, 4: 20, 5: 38, 6: 42, 7: 20,
            8: 11, 9: 24, 10: 14, 11: 40, 12: 55, 13: 55,
            14: 42, 15: 42, 16: 45, 17: 45, 18: 17, 19: 18,
            20: 20, 21: 16, 22: 65, 23: 65,
        }
        for index, width in widths.items():
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet["H"][1:]:
            cell.number_format = "0"
        for cell in sheet["R"][1:]:
            cell.number_format = "0.00"

        XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
        temporary = XLSX_PATH.with_suffix(".tmp.xlsx")
        workbook.save(temporary)
        os.replace(temporary, XLSX_PATH)
        return XLSX_PATH


def result_count() -> int:
    with _lock:
        if not CSV_PATH.exists():
            return 0
        with CSV_PATH.open("r", newline="", encoding="utf-8-sig") as handle:
            return max(0, sum(1 for _ in csv.reader(handle)) - 1)
